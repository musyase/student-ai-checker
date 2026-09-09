"""All model responses here are explicit HTTP mocks, not real AI grading."""
from __future__ import annotations

import base64
import io
import json
import logging
import struct
import zlib

import httpx
import pytest
from docx import Document

from student_ai_checker import cli
from student_ai_checker.ai import OllamaClient
from student_ai_checker.documents import DocxExtractor
from student_ai_checker.documents.template import compare_template
from student_ai_checker.exceptions import OllamaResponseError
from student_ai_checker.grading import StudentWorkChecker
from student_ai_checker.grading.approval import criteria_digest
from student_ai_checker.grading.per_criterion import PerCriterionChecker
from student_ai_checker.grading.template_checker import TemplateWorkChecker
from student_ai_checker.models import CriteriaSet, Criterion, GradingResponse


def png(red: int) -> bytes:
    """Valid small synthetic PNG; never a real student's image."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data))
    header = struct.pack('!2I5B', 8, 8, 8, 2, 0, 0, 0)
    rows = (b'\0' + bytes([red, 80, 100]) * 8) * 8
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', header) + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b'')


@pytest.fixture
def work(tmp_path):
    template = Document()
    template.add_paragraph('INHERITED_INSTRUCTION')
    template.add_paragraph('Ответ: [Заполните]')
    template_path = tmp_path / 'Бланк.docx'
    template.save(template_path)
    submission = Document(template_path)
    submission.paragraphs[1].text = 'Ответ: [PRIVATE_STUDENT_ANSWER]'
    for red in (10, 230):
        submission.add_picture(io.BytesIO(png(red)))
    submission.add_table(rows=1, cols=1).cell(0, 0).text = 'PRIVATE_TABLE_ANSWER'
    submission_path = tmp_path / 'Работа.docx'
    submission.save(submission_path)
    extracted = DocxExtractor().extract(submission_path)
    comparison = compare_template(extracted, DocxExtractor().extract(template_path))
    criteria = CriteriaSet(version=7, criteria=[
        Criterion(id=f'c{i}', title=f'Criterion {i}', description=f'ONLY_TASK_{i}',
                  max_score=2, allowed_scores=[0, 1, 2],
                  evidence_requirements=[{'description': 'Image or text', 'kinds': ['image', 'paragraph']}])
        for i in (1, 2)
    ])
    return submission_path, template_path, extracted, comparison, criteria


def answer(c: dict, *, status='completed', score=2, evidence=None) -> dict:
    manual = status in ('insufficient_evidence', 'unable_to_check')
    return {
        'criterion_id': c['id'], 'title': c['title'], 'max_score': c['max_score'],
        'proposed_score': None if manual else score, 'status': status,
        'explanation': 'MOCK_ONLY_DO_NOT_REUSE_AS_PROMPT',
        'evidence_ids': evidence or ['DOC1:IMG:0002'], 'student_feedback': '',
        'manual_review_required': manual, 'manual_review_reason': 'MOCK uncertainty' if manual else None,
    }


def mock_transport(responder=None):
    """Capture actual HTTP-client payloads; substitute only the local server."""
    posts = []
    def handler(request):
        if request.url.path == '/api/tags':
            return httpx.Response(200, json={'models': [{'name': 'qwen3.5:4b'}]})
        assert request.url.path == '/api/chat'
        payload = json.loads(request.content)
        posts.append(payload)
        criteria_json = payload['messages'][0]['content'].split('APPROVED_CRITERIA:\n', 1)[1]
        current, _ = json.JSONDecoder().raw_decode(criteria_json)
        if responder:
            special = responder(payload, current, len(posts))
            if isinstance(special, httpx.Response):
                return special
            content = special
        else:
            content = json.dumps({'results': [answer(c) for c in current['criteria']]})
        return httpx.Response(200, json={'message': {'content': content}, 'done': True, 'done_reason': 'stop'})
    http = httpx.Client(transport=httpx.MockTransport(handler))
    local = OllamaClient(base_url='http://127.0.0.1:11434', model='qwen3.5:4b', client=http)
    return http, local, posts


@pytest.mark.parametrize('template_mode', [False, True])
def test_single_criterion_scope_preserves_all_evidence_and_images(work, template_mode, caplog):
    _, _, document, comparison, criteria = work
    before = document.model_dump_json(), criteria.model_dump_json()
    expected_doc = comparison.filtered_document if template_mode else document
    http, local, posts = mock_transport()
    caplog.set_level(logging.INFO)
    with http:
        checker = TemplateWorkChecker(local, comparison) if template_mode else StudentWorkChecker(local)
        result = PerCriterionChecker(checker).check(document, criteria)
    assert len(posts) == 2
    assert [r.criterion_id for r in result.results] == ['c1', 'c2']
    assert before == (document.model_dump_json(), criteria.model_dump_json())
    for index, payload in enumerate(posts, 1):
        system, user = payload['messages']
        assert f'ONLY_TASK_{index}' in system['content']
        assert f'ONLY_TASK_{3-index}' not in system['content']
        assert 'MOCK_ONLY_DO_NOT_REUSE_AS_PROMPT' not in str(payload)
        assert 'PRIVATE_STUDENT_ANSWER' not in system['content']
        assert 'PRIVATE_STUDENT_ANSWER' in user['content']
        assert 'PRIVATE_TABLE_ANSWER' in user['content']
        assert ('INHERITED_INSTRUCTION' in user['content']) is (not template_mode)
        data = json.loads(user['content'].split('\n', 1)[1])
        assert {e['evidence_id'] for e in data['untrusted_student_evidence']} == expected_doc.evidence_ids
        assert [e['evidence_id'] for e in data['attached_images']] == [e.id for e in expected_doc.images]
        assert [base64.b64decode(b) for b in user['images']] == [e.binary_data for e in document.images]
        assert payload['format'] == GradingResponse.model_json_schema()
        assert payload['options'] == {'temperature': 0, 'num_ctx': 16384}
        assert payload['model'] == 'qwen3.5:4b' and payload['think'] is False
    assert 'Проверка критерия 1/2' in caplog.text and 'Проверка критерия 2/2' in caplog.text
    assert 'PRIVATE_' not in caplog.text


@pytest.mark.parametrize('field,value', [
    ('criterion_id', 'c2'), ('title', 'AI changed title'), ('max_score', 50),
    ('proposed_score', 0.5), ('status', 'invented'), ('evidence_ids', ['DOC1:IMG:9999']),
    ('evidence_ids', ['DOC1:P:0001']),  # A real but inherited template instruction.
    ('evidence_ids', ['DOC1:META:0001']),  # Real metadata cannot justify positive credit.
])
def test_bad_first_answer_repairs_only_current_criterion(work, field, value):
    _, _, doc, comparison, criteria = work
    def respond(payload, current, call):
        result = answer(current['criteria'][0])
        if call == 1:
            result[field] = value
        return json.dumps({'results': [result]})
    http, local, posts = mock_transport(respond)
    with http:
        result = PerCriterionChecker(TemplateWorkChecker(local, comparison)).check(doc, criteria)
    assert len(posts) == 3
    assert 'ONLY_TASK_1' in posts[1]['messages'][0]['content']
    assert 'REPAIR_NOTICE' in posts[1]['messages'][1]['content']
    assert 'REPAIR_NOTICE' not in posts[2]['messages'][1]['content']
    assert [r.proposed_score for r in result.results] == [2, 2]


@pytest.mark.parametrize('status', ['insufficient_evidence', 'unable_to_check'])
def test_valid_uncertainty_is_not_retried_or_changed_to_points(work, status):
    _, _, doc, comparison, criteria = work
    def respond(payload, current, call):
        return json.dumps({'results': [answer(current['criteria'][0], status=status if call == 1 else 'completed')]})
    http, local, posts = mock_transport(respond)
    with http:
        result = PerCriterionChecker(TemplateWorkChecker(local, comparison)).check(doc, criteria)
    summary = cli._score_summary(criteria, result)
    assert len(posts) == 2 and result.results[0].proposed_score is None
    assert summary['proposed_total'] is None and summary['scored_points'] == 2
    assert summary['unresolved_count'] == 1 and summary['teacher_confirmed'] is False
    assert summary['grade'] is None


@pytest.mark.parametrize('payload', ['not JSON', '{"results": []}'])
def test_second_validation_failure_aborts_without_fabrication(work, payload):
    _, _, doc, comparison, criteria = work
    http, local, posts = mock_transport(lambda *args: payload)
    with http, pytest.raises(OllamaResponseError):
        PerCriterionChecker(TemplateWorkChecker(local, comparison)).check(doc, criteria)
    assert len(posts) == 2
    assert all('ONLY_TASK_1' in p['messages'][0]['content'] for p in posts)


@pytest.mark.parametrize('code', [400, 500])
def test_transport_failure_is_not_retried_or_converted_to_a_grade(work, code):
    _, _, doc, comparison, criteria = work
    http, local, posts = mock_transport(lambda *args: httpx.Response(code, json={'error': 'MOCK server failure'}))
    with http, pytest.raises(OllamaResponseError):
        PerCriterionChecker(TemplateWorkChecker(local, comparison)).check(doc, criteria)
    assert len(posts) == 1


def test_cancel_stops_before_next_criterion(work):
    _, _, doc, comparison, criteria = work
    def cancel(*args):
        raise KeyboardInterrupt
    http, local, posts = mock_transport(cancel)
    with http, pytest.raises(KeyboardInterrupt):
        PerCriterionChecker(TemplateWorkChecker(local, comparison)).check(doc, criteria)
    assert len(posts) == 1


def prepare_cli_files(work, tmp_path):
    source, template, _, comparison, criteria = work
    rubric = tmp_path / 'Критерии.json'
    rubric.write_text(criteria.model_dump_json(), encoding='utf-8')
    approval = tmp_path / 'Утверждение.json'
    approval.write_text(json.dumps({
        'status': 'approved', 'approved_at': '2026-09-09T19:15:26Z', 'approved_by': 'teacher',
        'digest_format': 'criteria-model-json-v1', 'criteria_version': criteria.version,
        'criteria_sha256': criteria_digest(criteria), 'template_sha256': comparison.template_sha256,
        'grade_scale': None,
    }), encoding='utf-8')
    output = tmp_path / 'Новая проверка.json'
    args = ['check', '--document', str(source), '--template', str(template), '--criteria', str(rubric),
            '--approval', str(approval), '--output', str(output)]
    return args, output, rubric, approval


@pytest.mark.parametrize('per_criterion', [False, True])
def test_cli_real_documents_mock_http_save_full_rubric_and_mode(work, tmp_path, monkeypatch, per_criterion):
    args, output, rubric, approval = prepare_cli_files(work, tmp_path)
    if per_criterion:
        args.append('--per-criterion')
    originals = {p: p.read_bytes() for p in [work[0], work[1], rubric, approval]}
    http, local, posts = mock_transport()
    monkeypatch.setattr(cli, 'OllamaClient', lambda **kwargs: local)
    with http:
        assert cli.main(args) == 0
    result = json.loads(output.read_text(encoding='utf-8'))
    assert len(posts) == (2 if per_criterion else 1)
    assert result['inference_options']['criterion_mode'] == ('per_criterion' if per_criterion else 'all_at_once')
    assert result['criteria_sha256'] == criteria_digest(work[4])
    assert result['approval']['criteria_version'] == 7
    assert result['summary']['max_score'] == 4 and result['summary']['proposed_total'] == 4
    assert result['summary']['teacher_confirmed'] is False and result['summary']['grade'] is None
    assert [r['criterion_id'] for r in result['results']] == ['c1', 'c2']
    assert originals == {p: p.read_bytes() for p in originals}


@pytest.mark.parametrize('failure', ['rubric_changed', 'template_changed', 'output_exists', 'missing_approval'])
def test_cli_preflight_stops_before_any_model_request(work, tmp_path, monkeypatch, failure):
    args, output, rubric, approval = prepare_cli_files(work, tmp_path)
    args.append('--per-criterion')
    if failure == 'rubric_changed':
        changed = json.loads(rubric.read_text())
        changed['criteria'][0]['description'] += 'changed'
        rubric.write_text(json.dumps(changed), encoding='utf-8')
    elif failure == 'template_changed':
        d = Document(work[1]); d.add_paragraph('New task'); d.save(work[1])
    elif failure == 'output_exists':
        output.write_text('OLD_CONFIRMED_RESULT', encoding='utf-8')
    else:
        approval.unlink()
    def unexpected_client(**kwargs):
        pytest.fail('Model client must not be constructed before preflight passes')
    monkeypatch.setattr(cli, 'OllamaClient', unexpected_client)
    assert cli.main(args) == 2
    if failure == 'output_exists':
        assert output.read_text() == 'OLD_CONFIRMED_RESULT'
    else:
        assert not output.exists()


def test_cli_second_criterion_http_error_does_not_save_partial_run(work, tmp_path, monkeypatch):
    args, output, _, _ = prepare_cli_files(work, tmp_path)
    args.append('--per-criterion')
    def respond(payload, current, call):
        if call == 2:
            return httpx.Response(500, json={'error': 'MOCK fail'})
        return json.dumps({'results': [answer(current['criteria'][0])]})
    http, local, posts = mock_transport(respond)
    monkeypatch.setattr(cli, 'OllamaClient', lambda **kwargs: local)
    with http:
        assert cli.main(args) == 2
    assert len(posts) == 2 and not output.exists()


def test_ambiguity_protection_remains_active(work):
    _, _, doc, comparison, criteria = work
    comparison.origins['DOC1:P:0002'] = 'ambiguous'
    http, local, posts = mock_transport()
    with http, pytest.raises(OllamaResponseError):
        PerCriterionChecker(TemplateWorkChecker(local, comparison)).check(doc, criteria)
    assert len(posts) == 2
