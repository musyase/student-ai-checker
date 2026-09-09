"""Opt-in semantic check against REAL local Ollama, including a negative control.

These are synthetic coloured images, not student work. A response containing only
null scores must not pass. Passing these cases does not establish grading quality.
"""
from __future__ import annotations

import os
import struct
import zlib

import pytest
from docx import Document

from student_ai_checker.ai import OllamaClient
from student_ai_checker.config import Settings
from student_ai_checker.documents import DocxExtractor
from student_ai_checker.documents.template import compare_template
from student_ai_checker.grading.visual_review import VisualReviewChecker
from student_ai_checker.models import CriteriaSet, Criterion

pytestmark = [
    pytest.mark.integration,
    pytest.mark.skipif(os.getenv("STUDENT_AI_CHECKER_RUN_OLLAMA") != "1",
                       reason="Requires explicit opt-in and real local Ollama"),
]


def _solid_png(rgb: tuple[int, int, int]) -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack("!I", len(data)) + kind + data + struct.pack("!I", zlib.crc32(kind + data) & 0xffffffff)
    header = struct.pack("!2I5B", 256, 192, 8, 2, 0, 0, 0)
    pixels = b"".join(b"\0" + bytes(rgb) * 256 for _ in range(192))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(pixels)) + chunk(b"IEND", b"")


@pytest.mark.parametrize("colour,expected", [("красным", 1), ("зелёным", 0)])
def test_real_visual_review_present_and_absent_feature(tmp_path, colour, expected):
    blank = Document()
    blank.add_paragraph("Задание: приложите изображения. Требования задаются отдельно.")
    blank_path = tmp_path / "template.docx"
    blank.save(blank_path)
    for index, rgb in enumerate(((255, 0, 0), (0, 0, 255))):
        path = tmp_path / f"image-{index}.png"
        path.write_bytes(_solid_png(rgb))
        blank.add_picture(str(path))
    work_path = tmp_path / "work.docx"
    blank.save(work_path)
    extractor = DocxExtractor()
    doc = extractor.extract(work_path)
    comp = compare_template(doc, extractor.extract(blank_path))
    rubric = CriteriaSet(criteria=[Criterion(
        id="colour-check", title="Проверка цвета изображения",
        description=f"1 балл, если есть изображение, полностью залитое {colour} цветом. "
                    "0 баллов, если все изображения хорошо видны, но такого цвета нет. "
                    "Прочие цвета не засчитываются. По тексту и имени файла цвет не определять.",
        max_score=1, allowed_scores=[0, 1],
        evidence_requirements=[{"description": "Цвет самого изображения", "kinds": ["image"]}],
    )])
    settings = Settings.from_env()
    with OllamaClient(base_url=settings.ollama_base_url, model=settings.ollama_model,
                      timeout_seconds=settings.ollama_timeout_seconds,
                      num_ctx=settings.ollama_num_ctx) as client:
        checker = VisualReviewChecker(client, comp)
        result = checker.check(doc, rubric).results[0]
    assert result.proposed_score == expected, result.model_dump_json(indent=2)
    if expected:
        assert doc.images[0].id in result.evidence_ids
    assert len(checker.audit[0]["image_reviews"]) == 2
    assert all(r["finding"] is not None for r in checker.audit[0]["image_reviews"])
