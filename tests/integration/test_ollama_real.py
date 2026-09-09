from __future__ import annotations

import base64
import os
from pathlib import Path

import pytest
from docx import Document

from student_ai_checker.ai import OllamaClient
from student_ai_checker.config import Settings
from student_ai_checker.documents import DocxExtractor
from student_ai_checker.grading import StudentWorkChecker
from student_ai_checker.models import CriteriaSet, Criterion, EvidenceRequirement

PNG_1X1 = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII=")

pytestmark = pytest.mark.integration


@pytest.mark.skipif(os.getenv("STUDENT_AI_CHECKER_RUN_OLLAMA") != "1", reason="set STUDENT_AI_CHECKER_RUN_OLLAMA=1 to run against real local Ollama")
def test_real_docx_to_ollama_structured_result(tmp_path: Path) -> None:
    image = tmp_path / "screen.png"
    image.write_bytes(PNG_1X1)
    source = tmp_path / "student.docx"
    doc = Document()
    doc.add_paragraph("Проверочный отчёт. Приложено изображение результата.")
    doc.add_picture(str(image))
    doc.save(source)
    criteria = CriteriaSet(criteria=[Criterion(id="image-present", title="Приложено изображение", description="В отчёте есть встроенное изображение.", max_score=1, allowed_scores=[0, 1], evidence_requirements=[EvidenceRequirement(description="Встроенное изображение", kinds=["image"])])])
    extracted = DocxExtractor().extract(source)
    settings = Settings.from_env()
    with OllamaClient(base_url=settings.ollama_base_url, model=settings.ollama_model, timeout_seconds=settings.ollama_timeout_seconds, num_ctx=settings.ollama_num_ctx) as ollama:
        result = StudentWorkChecker(ollama).check(extracted, criteria)
    assert result.results[0].criterion_id == "image-present"
    assert set(result.results[0].evidence_ids) <= extracted.evidence_ids
