from __future__ import annotations

import base64
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
)

pytestmark = pytest.mark.integration


@pytest.mark.skipif(
    os.getenv("STUDENT_AI_CHECKER_RUN_OLLAMA") != "1",
    reason="set STUDENT_AI_CHECKER_RUN_OLLAMA=1 to run against real local Ollama",
)
def test_real_cli_docx_to_result_json(tmp_path: Path) -> None:
    """Exercise the complete Stage 1 process through the real CLI and Ollama."""

    image_path = tmp_path / "screen.png"
    image_path.write_bytes(PNG_1X1)

    document_path = tmp_path / "student.docx"
    document = Document()
    document.add_paragraph("Проверочный отчёт. Приложено изображение результата.")
    document.add_picture(str(image_path))
    document.save(document_path)

    criteria_path = tmp_path / "criteria.json"
    criteria_path.write_text(
        json.dumps(
            {
                "version": 1,
                "criteria": [
                    {
                        "id": "image-present",
                        "title": "Приложено изображение",
                        "description": "В отчёте есть встроенное изображение.",
                        "max_score": 1,
                        "allowed_scores": [0, 1],
                        "evidence_requirements": [
                            {
                                "description": "Встроенное изображение",
                                "kinds": ["image"],
                                "required": True,
                            }
                        ],
                    }
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    result_path = tmp_path / "result.json"
    environment = os.environ.copy()
    environment["PYTHONUTF8"] = "1"

    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "student_ai_checker.cli",
            "check",
            "--document",
            str(document_path),
            "--criteria",
            str(criteria_path),
            "--output",
            str(result_path),
        ],
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=360,
        check=False,
    )

    assert completed.returncode == 0, (
        f"CLI failed with code {completed.returncode}\n"
        f"STDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}"
    )
    assert result_path.is_file()

    payload = json.loads(result_path.read_text(encoding="utf-8"))
    assert payload["schema_version"] == 1
    assert payload["document"]["source_name"] == "student.docx"
    assert payload["document"]["source_sha256"]
    assert payload["results"][0]["criterion_id"] == "image-present"
    assert payload["results"][0]["status_label_ru"]

    known_evidence_ids = {
        evidence["id"] for evidence in payload["resolved_evidence"]
    }
    assert set(payload["results"][0]["evidence_ids"]) <= known_evidence_ids
