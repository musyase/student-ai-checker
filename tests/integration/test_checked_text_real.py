"""Opt-in REAL local model tests of text meaning, not HTTP mocks or student grades."""
from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document

from student_ai_checker.ai import OllamaClient
from student_ai_checker.config import Settings
from student_ai_checker.documents import DocxExtractor
from student_ai_checker.documents.template import compare_template
from student_ai_checker.grading.checked_text import CHECKED_SCORING_VERSION, PARTS
from student_ai_checker.grading.referenced_text import check_text_references as check_text_components
from student_ai_checker.models import CriteriaSet

pytestmark = [pytest.mark.integration, pytest.mark.skipif(
    os.getenv('STUDENT_AI_CHECKER_RUN_OLLAMA') != '1', reason='Requires real local Ollama and explicit opt-in')]

# Deliberately synthetic answers, written from the approved assignment, not copied
# from a student's report. Expected scores are used ONLY by the assertions below.
CASES = [
    ('add_full', 'pr1.add-button', 'После нажатия: [Задача «Подготовиться к паре» добавится к двум предыдущим. Их станет три. Поле предлагаю очищать.]', 3),
    ('add_partial', 'pr1.add-button', 'После нажатия: [Задача «Подготовиться к паре» добавится в список. Всего будет три задачи. Что делать с полем ввода, не указано.]', 2),
    ('add_wrong', 'pr1.add-button', 'После нажатия: [Новая задача не добавится. В списке останется две задачи. Текст в поле ввода предлагаю стирать.]', 1),
    ('add_blank', 'pr1.add-button', None, 0),
    ('usability_full', 'pr1.usability', ['[Назвать кнопку «Добавить задачу».]', '[Подписать поле «Новая задача».]', '[Сделать шрифт крупнее.]'], 3),
    ('usability_wrong', 'pr1.usability', ['[Оставить надпись «ОК».]', '[Не добавлять подпись и подсказку.]', '[Сделать текст ещё мельче.]'], 0),
    ('installation_full', 'pr1.installation-report', 'Ход работы: [Среда уже была установлена. Я нашёл её через меню и открыл. Появилось стартовое окно. Затруднений не было.]', 3),
    ('installation_partial', 'pr1.installation-report', 'Ход работы: [На компьютере программа уже установлена.]', 1),
]


def write_text_control_report(directory: Path, cases: list[dict]) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    name = 'text-controls-' + datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S') + '-' + uuid4().hex[:8] + '.json'
    path = directory/name
    with path.open('x', encoding='utf-8') as stream:
        json.dump({'purpose':'synthetic_text_controls_not_student_grading',
                   'scoring_version':CHECKED_SCORING_VERSION, 'model_output_is_untrusted':True,
                   'cases':cases}, stream, ensure_ascii=False, indent=2, allow_nan=False)
    return path


@pytest.fixture(scope='module')
def control_audit(tmp_path_factory):
    configured = os.getenv('STUDENT_AI_CHECKER_TEST_ARTIFACT_DIR')
    directory = Path(configured) if configured else tmp_path_factory.mktemp('text-controls')
    cases = []
    try:
        yield cases
    finally:
        print('\nКонтрольный журнал текста:', write_text_control_report(directory, cases))


@pytest.mark.parametrize('case_id,criterion_id,answer,expected', CASES, ids=[c[0] for c in CASES])
def test_real_checked_text(tmp_path, case_id, criterion_id, answer, expected, control_audit):
    cs = CriteriaSet.model_validate_json((Path(__file__).parents[2]/'examples/pr1.criteria.approved.json').read_text(encoding='utf-8'))
    c = cs.by_id()[criterion_id]
    blank = Document()
    blank.add_paragraph(c.description)
    if criterion_id == 'pr1.usability':
        table = blank.add_table(rows=4, cols=2)
        table.cell(0,0).text='Что неудобно'; table.cell(0,1).text='Как исправить'
        for index, problem in enumerate(['Кнопка добавления подписана «ОК».', 'Нет подписи или подсказки у поля.', 'Текст слишком мелкий.'], 1):
            table.cell(index,0).text=problem
            table.cell(index,1).text='[Ваше исправление.]'
    else:
        label = 'После нажатия' if criterion_id == 'pr1.add-button' else 'Ход работы'
        blank.add_paragraph(label+': [Напишите здесь.]')
    template_path=tmp_path/'template.docx'; blank.save(template_path)
    if isinstance(answer, list):
        for i,text in enumerate(answer,1):table.cell(i,1).text=text
    elif answer is not None:
        blank.paragraphs[1].text=answer
    work_path=tmp_path/'work.docx'; blank.save(work_path)
    extractor=DocxExtractor(); source=extractor.extract(work_path)
    comparison=compare_template(source,extractor.extract(template_path))
    settings=Settings.from_env()
    record={'case':case_id, 'expected_score':expected, 'model':settings.ollama_model,
            'num_ctx':settings.ollama_num_ctx, 'source_sha256':source.source_sha256,
            'result':None, 'audit':None, 'passed':False}
    control_audit.append(record)
    with OllamaClient(base_url=settings.ollama_base_url, model=settings.ollama_model,
                      timeout_seconds=settings.ollama_timeout_seconds, num_ctx=settings.ollama_num_ctx) as client:
        client.assert_ready()
        result,audit=check_text_components(client,c,PARTS[c.id],comparison.filtered_document)
    record.update(result=result.model_dump(mode='json'),audit=audit)
    # null must fail; quotes and exact fixed component IDs were already validated.
    assert result.proposed_score == expected, json.dumps(record, ensure_ascii=False, indent=2)
    assert_component_results(record, EXPECTED_COMPONENTS[case_id])
    record['passed']=True


# Assertions only: never pass these expected verdicts or scores to the model.
EXPECTED_COMPONENTS = {
    'add_full': ('met', 'met', 'met'),
    'add_partial': ('met', 'met', 'not_met'),
    'add_wrong': ('not_met', 'not_met', 'met'),
    'add_blank': ('not_met', 'not_met', 'not_met'),
    'usability_full': ('met', 'met', 'met'),
    'usability_wrong': ('not_met', 'not_met', 'not_met'),
    'installation_full': ('met', 'met', 'met'),
    'installation_partial': ('met', 'not_met', 'not_met'),
    'fixes_paraphrased_wrong': ('not_met', 'not_met', 'not_met'),
    'fixes_paraphrased_full': ('met', 'met', 'met'),
    'fixes_mixed': ('met', 'not_met', 'met'),
    'installation_initial_only': ('met', 'not_met', 'not_met'),
    'installation_no_problem_sentence': ('met', 'met', 'met'),
    'installation_interrupted_described': ('met', 'met', 'met'),
    'add_keep_input': ('met', 'met', 'met'),
    'add_only_input': ('not_met', 'not_met', 'met'),
}

# New contrastive development controls, not an independent acceptance dataset.
# They vary wording, negation, partial content, and conditional requirements.
VARIANT_CASES = [
    ('fixes_paraphrased_wrong', 'pr1.usability',
     ['[Не менять название кнопки.]', '[Поле оставим без подписи и пояснения назначения.]',
      '[Уменьшить кегль надписей.]'], 0),
    ('fixes_paraphrased_full', 'pr1.usability',
     ['[Назвать кнопку «Создать задачу», чтобы было понятно действие.]',
      '[Поставить рядом с полем подпись «Название задачи».]',
      '[Сделать буквы достаточно большими для чтения.]'], 3),
    ('fixes_mixed', 'pr1.usability',
     ['[Подписать кнопку «Добавить запись».]',
      '[Скрыть подсказку и оставить поле без подписи.]', '[Увеличить буквы.]'], 2),
    ('installation_initial_only', 'pr1.installation-report',
     'Ход работы: [До занятия среда была на этом ПК. Больше ничего о ходе работы не записано.]', 1),
    ('installation_no_problem_sentence', 'pr1.installation-report',
     'Ход работы: [На ПК уже имелась среда. Я запустил её через ярлык. Открылось начальное окно.]', 3),
    ('installation_interrupted_described', 'pr1.installation-report',
     'Ход работы: [Редактора на компьютере не было. Я запустил установщик. Он остановился на запросе прав. Я прекратил установку и попросил преподавателя помочь.]', 3),
    ('add_keep_input', 'pr1.add-button',
     'После нажатия: [Новая запись появится под прежними двумя. В списке будет три задачи. Введённый текст предлагаю оставить, чтобы исправить его перед следующим добавлением.]', 3),
    ('add_only_input', 'pr1.add-button',
     'После нажатия: [Не описываю изменение списка и количества задач. Предлагаю освобождать поле от введённого текста.]', 1),
]


def assert_component_results(record: dict, expected_verdicts: tuple[str, ...]) -> None:
    """A correct sum cannot conceal two compensating component mistakes."""
    record['expected_components'] = list(expected_verdicts)
    audit = record['audit']
    rows = audit['components']
    expected_ids = [part.id for part in PARTS[record['result']['criterion_id']]]
    actual_ids = [row['check_id'] for row in rows]
    states = tuple(row['verdict'] for row in rows)
    points = [row['points'] for row in rows]
    detail = json.dumps(record, ensure_ascii=False, indent=2)
    assert actual_ids == expected_ids, detail
    assert states == expected_verdicts, detail
    assert points == [int(state == 'met') for state in expected_verdicts], detail


@pytest.mark.parametrize('case_id,criterion_id,answer,expected', VARIANT_CASES,
                         ids=[c[0] for c in VARIANT_CASES])
def test_real_checked_text_variants(tmp_path, case_id, criterion_id, answer, expected, control_audit):
    # Reuse the exact document/extraction/reference path, not fabricated observations.
    test_real_checked_text(tmp_path, case_id, criterion_id, answer, expected, control_audit)
