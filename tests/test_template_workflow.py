"""Synthetic fixtures only. OllamaStub is a unit-test mock, never production AI."""
from __future__ import annotations

import base64
import json
from pathlib import Path

import pytest
from docx import Document

from student_ai_checker import cli
from student_ai_checker.documents import DocxExtractor
from student_ai_checker.documents.template import compare_template
from student_ai_checker.exceptions import DocumentExtractionError, GradingValidationError, OllamaResponseError, StudentAICheckerError
from student_ai_checker.grading.approval import criteria_digest, verify_approval
from student_ai_checker.grading.template_checker import TemplateGradingValidator, TemplateWorkChecker
from student_ai_checker.models import CriteriaSet, Criterion, CriterionResult, EvidenceItem, EvidenceKind, ExtractedDocument, ExtractionIssue, ExtractionIssueSeverity, GradingResponse


@pytest.fixture
def criteria():
    return CriteriaSet(criteria=[Criterion(id="c1", title="Название", description="Названо приложение.", max_score=1, allowed_scores=[0, 1])])


@pytest.fixture
def pair(tmp_path):
    blank = Document()
    blank.add_paragraph("Задание: укажите название знакомого приложения, которым вы пользуетесь.")
    blank.add_paragraph("Название: [Напишите здесь.]")
    table = blank.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Проблема"
    table.cell(0, 1).text = "Ответ"
    table.cell(1, 0).text = "Мелкий текст"
    table.cell(1, 1).text = "[Ваше исправление.]"
    blank_path = tmp_path / "Бланк.docx"
    blank.save(blank_path)
    blank.paragraphs[1].text = "Название: [Приложение А]"
    table.cell(1, 1).text = "[Увеличить шрифт]"
    work_path = tmp_path / "Работа.docx"
    blank.save(work_path)
    extractor = DocxExtractor()
    return blank_path, work_path, extractor.extract(blank_path), extractor.extract(work_path)


def approval_file(tmp_path, criteria, template_hash):
    path = tmp_path / "approval.json"
    path.write_text(json.dumps({
        "status": "approved", "approved_at": "2026-09-09T19:15:26Z", "approved_by": "teacher",
        "digest_format": "criteria-model-json-v1", "criteria_version": criteria.version,
        "criteria_sha256": criteria_digest(criteria), "template_sha256": template_hash, "grade_scale": None,
    }), encoding="utf-8")
    return path


def response(evidence_id="DOC1:P:0002", score=1, status="completed"):
    manual = score is None
    return GradingResponse(results=[CriterionResult(
        criterion_id="c1", title="Название", max_score=1, proposed_score=score, status=status,
        explanation="Описание по указанному фрагменту.", evidence_ids=[evidence_id],
        manual_review_required=manual, manual_review_reason="Нужно уточнение." if manual else None,
    )])


class OllamaStub:
    """Unit-only transport double; no production fallback uses this class."""
    def __init__(self, replies):
        self.replies = iter(replies)
        self.calls = []
    def assert_ready(self):
        pass
    def chat_structured(self, **kwargs):
        self.calls.append(kwargs)
        return next(self.replies)
    def __enter__(self):
        return self
    def __exit__(self, *args):
        pass


def test_pr1_snapshot_is_exactly_the_approved_draft():
    root = Path(__file__).parents[1]
    blank = root / "examples/pr1.criteria.draft.json"
    approved = root / "examples/pr1.criteria.approved.json"
    assert approved.read_bytes() == blank.read_bytes()
    c = CriteriaSet.model_validate_json(approved.read_text(encoding="utf-8"))
    assert sum(x.max_score for x in c.criteria) == 25
    record = verify_approval(root / "examples/pr1.approval.json", criteria=c,
                             template_sha256="c2193a2dccad9fbd04837e3f5a899cf07999b1734cc1dcaf4166779a7bd89956")
    assert record.grade_scale is None


@pytest.mark.parametrize("field,value", [("status", "draft"), ("approved_by", "AI"), ("grade_scale", {"90": 5}), ("criteria_version", 2), ("template_sha256", "0" * 64)])
def test_invalid_approval_is_rejected(tmp_path, pair, criteria, field, value):
    p = approval_file(tmp_path, criteria, pair[2].source_sha256)
    data = json.loads(p.read_text()); data[field] = value
    p.write_text(json.dumps(data))
    with pytest.raises(StudentAICheckerError):
        verify_approval(p, criteria=criteria, template_sha256=pair[2].source_sha256)


def test_criteria_mutation_after_approval_is_rejected(tmp_path, pair, criteria):
    path = approval_file(tmp_path, criteria, pair[2].source_sha256)
    criteria.criteria[0].description = "Другое требование"
    with pytest.raises(StudentAICheckerError, match="Критерии изменились"):
        verify_approval(path, criteria=criteria, template_sha256=pair[2].source_sha256)


def test_criteria_digest_ignores_formatting_and_numeric_spelling(criteria):
    text = criteria.model_dump_json(indent=4).replace("1.0", "1")
    assert criteria_digest(CriteriaSet.model_validate_json(text)) == criteria_digest(criteria)


def test_comparison_preserves_answers_brackets_tables_original_ids(pair):
    before = pair[3].model_dump()
    comp = compare_template(pair[3], pair[2])
    assert pair[3].model_dump() == before
    assert comp.origins["DOC1:P:0001"] == "template_match"
    assert comp.origins["DOC1:P:0002"] == "candidate"
    assert comp.origins["DOC1:T:0001:R:0002:C:0002"] == "candidate"
    assert comp.origins["DOC1:T:0001:R:0002:C:0001"] == "template_match"
    assert comp.filtered_document.evidence_by_id()["DOC1:P:0002"].text == "Название: [Приложение А]"


def test_reordered_or_duplicated_instruction_still_excluded(pair):
    doc = pair[3].model_copy(deep=True)
    doc.evidence.append(EvidenceItem(id="DOC1:P:0999", kind="paragraph", order_index=999,
                                   text=pair[2].evidence[1].text))
    assert compare_template(doc, pair[2]).origins["DOC1:P:0999"] == "template_match"


def test_whitespace_is_not_a_student_answer(pair):
    doc = pair[2].model_copy(deep=True)
    for e in doc.evidence:
        if e.text:
            e.text = e.text.replace(" ", " \n\t ")
    comp = compare_template(doc, pair[2])
    assert not any(v == "candidate" for v in comp.origins.values())


@pytest.mark.parametrize("citation", ["DOC1:P:0001", "DOC1:META:0001"])
def test_blank_template_cannot_get_positive_points(pair, criteria, citation):
    comp = compare_template(pair[2], pair[2])
    with pytest.raises(GradingValidationError):
        TemplateGradingValidator(comp).validate(response(citation), criteria=criteria, document=comp.filtered_document)


@pytest.mark.parametrize("score,status", [(0, "not_completed"), (None, "insufficient_evidence")])
def test_blank_template_can_use_metadata_for_absence_or_uncertainty(pair, criteria, score, status):
    comp = compare_template(pair[2], pair[2])
    out = TemplateGradingValidator(comp).validate(response("DOC1:META:0001", score, status), criteria=criteria, document=comp.filtered_document)
    assert out.results[0].proposed_score == score


def test_changed_instruction_is_ambiguous_not_a_correct_answer(pair, criteria):
    doc = pair[3].model_copy(deep=True)
    doc.evidence[1].text += "!"
    comp = compare_template(doc, pair[2])
    assert comp.origins["DOC1:P:0001"] == "ambiguous"
    with pytest.raises(GradingValidationError, match="ambiguous"):
        TemplateGradingValidator(comp).validate(response(), criteria=criteria, document=comp.filtered_document)
    unresolved = response("DOC1:P:0002", None, "insufficient_evidence")
    TemplateGradingValidator(comp).validate(unresolved, criteria=criteria, document=comp.filtered_document)


@pytest.mark.parametrize("side", [2, 3])
def test_extraction_warning_stops_before_grading(pair, side):
    pair[side].issues.append(ExtractionIssue(code="test", message="Unsupported element", severity=ExtractionIssueSeverity.WARNING))
    with pytest.raises(DocumentExtractionError):
        compare_template(pair[3], pair[2])


def test_empty_template_is_not_accepted(pair):
    template = pair[2].model_copy(update={"evidence": []})
    with pytest.raises(DocumentExtractionError):
        compare_template(pair[3], template)


def test_image_bytes_not_filename_determine_inherited_image(pair):
    template = pair[2].model_copy(deep=True)
    document = pair[3].model_copy(deep=True)
    for doc, identifier, data in ((template, "DOC1:IMG:0001", b"same bytes"), (document, "DOC1:IMG:0099", b"same bytes"), (document, "DOC1:IMG:0100", b"different bytes")):
        doc.evidence.append(EvidenceItem(id=identifier, kind="image", order_index=200, binary_data=data))
    comp = compare_template(document, template)
    assert comp.origins["DOC1:IMG:0099"] == "template_match"
    assert comp.origins["DOC1:IMG:0100"] == "candidate"


def test_image_without_bytes_is_not_absence(pair):
    pair[3].evidence.append(EvidenceItem(id="DOC1:IMG:0001", kind="image", order_index=200))
    with pytest.raises(DocumentExtractionError):
        compare_template(pair[3], pair[2])


def test_duplicate_evidence_is_rejected(pair):
    pair[3].evidence.append(pair[3].evidence[1])
    with pytest.raises(DocumentExtractionError):
        compare_template(pair[3], pair[2])


def test_checker_filters_prompt_but_retries_invalid_evidence_once(pair, criteria):
    comp = compare_template(pair[3], pair[2])
    stub = OllamaStub([response("DOC1:P:0001").model_dump_json(), response().model_dump_json()])
    out = TemplateWorkChecker(stub, comp).check(pair[3], criteria)
    assert len(stub.calls) == 2
    assert "Задание: укажите" not in stub.calls[0]["user_prompt"]
    assert "Приложение А" in stub.calls[0]["user_prompt"]
    assert out.results[0].proposed_score == 1


def test_two_invalid_replies_do_not_create_fake_score(pair, criteria):
    comp = compare_template(pair[3], pair[2])
    bad = response("DOC1:META:0001").model_dump_json()
    stub = OllamaStub([bad, bad])
    with pytest.raises(OllamaResponseError):
        TemplateWorkChecker(stub, comp).check(pair[3], criteria)
    assert len(stub.calls) == 2


def test_real_image_bytes_are_forwarded_by_checker(pair, criteria):
    image = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII=")
    pair[3].evidence.append(EvidenceItem(id="DOC1:IMG:0001", kind="image", order_index=200, binary_data=image, mime_type="image/png"))
    stub = OllamaStub([response().model_dump_json()])
    TemplateWorkChecker(stub, compare_template(pair[3], pair[2])).check(pair[3], criteria)
    assert stub.calls[0]["images"] == [image]


def test_comparison_cannot_be_reused_for_another_submission(pair, criteria):
    changed = pair[3].model_copy(update={"source_sha256": "different"})
    stub = OllamaStub([])
    with pytest.raises(GradingValidationError):
        TemplateWorkChecker(stub, compare_template(pair[3], pair[2])).check(changed, criteria)
    assert stub.calls == []


def test_missing_score_is_not_silently_zero(criteria):
    summary = cli._score_summary(criteria, response("DOC1:META:0001", None, "unable_to_check"))
    assert summary["proposed_total"] is None
    assert summary["unresolved_count"] == 1
    assert summary["grade"] is None
    assert summary["teacher_confirmed"] is False


@pytest.mark.parametrize("mode", ["legacy", "template"])
def test_cli_creates_json_and_preserves_previous_result(tmp_path, pair, criteria, monkeypatch, mode):
    rubric = tmp_path / "criteria.json"
    rubric.write_text(criteria.model_dump_json(), encoding="utf-8")
    output = tmp_path / "results" / "проверка.json"
    stub = OllamaStub([response().model_dump_json()])
    monkeypatch.setattr(cli, "OllamaClient", lambda **kwargs: stub)
    args = ["check", "--document", str(pair[1]), "--criteria", str(rubric), "--output", str(output)]
    if mode == "template":
        approval = approval_file(tmp_path, criteria, pair[2].source_sha256)
        args += ["--template", str(pair[0]), "--approval", str(approval)]
    assert cli.main(args) == 0
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["summary"]["proposed_total"] == 1
    assert payload["summary"]["grade"] is None
    assert bool(payload["approval"]) == (mode == "template")
    original = output.read_bytes()
    assert cli.main(args) == 2
    assert len(stub.calls) == 1
    assert output.read_bytes() == original


def test_template_mode_requires_approval_before_io_or_ai(tmp_path):
    assert cli.main(["check", "--document", "missing.docx", "--criteria", "missing.json", "--template", "blank.docx", "--output", str(tmp_path/"result.json")]) == 2
