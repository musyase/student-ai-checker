from __future__ import annotations

import base64
from pathlib import Path

import pytest
from docx import Document

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
)


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    image_path = tmp_path / "скриншот.png"
    image_path.write_bytes(PNG_1X1)

    path = tmp_path / "Работа Иванов.docx"
    document = Document()
    document.add_paragraph("Установил Android Studio и создал проект.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Проблема"
    table.cell(0, 1).text = "SDK не найден — установил через SDK Manager."
    paragraph = document.add_paragraph("Результат:")
    paragraph.add_run().add_picture(str(image_path))
    document.save(path)
    return path
